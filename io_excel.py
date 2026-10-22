"""Excelテンプレの読込・結果出力・サンプル生成"""
from __future__ import annotations

import io
import random
from datetime import date, datetime, timedelta

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SHEET_STAFF = "スタッフ"
SHEET_PREF = "希望"
SHEET_REQ = "必要人数"

SHIFT_COLORS = ["FFE699", "9DC3E6", "C6E0B4", "F4B183"]  # 早番, 遅番, ...
WEEKDAY = "月火水木金土日"

HEAD_FILL = PatternFill("solid", fgColor="305496")
HEAD_FONT = Font(bold=True, color="FFFFFF")


def _to_date(v) -> date:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    return datetime.strptime(str(v)[:10].replace("/", "-"), "%Y-%m-%d").date()


class TemplateError(ValueError):
    """テンプレの書式エラー(利用者向けの日本語メッセージ)"""


def read_template(file) -> tuple[list, list, list, dict, dict]:
    """Excelを読み込み (staff, days, shifts, required, prefs) を返す。"""
    try:
        wb = load_workbook(file, data_only=True)
    except Exception as e:  # noqa: BLE001
        raise TemplateError(f"Excelファイルを開けませんでした: {e}") from e

    for name in (SHEET_STAFF, SHEET_PREF, SHEET_REQ):
        if name not in wb.sheetnames:
            raise TemplateError(f"シート「{name}」が見つかりません。テンプレートをご利用ください。")

    # --- スタッフ ---
    staff = []
    for row in wb[SHEET_STAFF].iter_rows(min_row=2, values_only=True):
        if not row or row[0] in (None, ""):
            continue
        try:
            max_days = int(row[1]) if row[1] not in (None, "") else 5
        except (TypeError, ValueError) as e:
            raise TemplateError(f"スタッフ「{row[0]}」の週上限日数が数字ではありません。") from e
        staff.append({
            "name": str(row[0]).strip(),
            "max_days": max_days,
            "is_leader": str(row[2]).strip() in ("◯", "○", "1", "TRUE", "True", "リーダー"),
        })
    if not staff:
        raise TemplateError("「スタッフ」シートに行がありません。")
    names = [s["name"] for s in staff]
    if len(set(names)) != len(names):
        raise TemplateError("「スタッフ」シートに同じ名前が重複しています。")

    # --- 必要人数 ---
    ws = wb[SHEET_REQ]
    header = [c.value for c in ws[1]]
    shifts = [str(h).strip() for h in header[1:] if h not in (None, "")]
    if len(shifts) < 1:
        raise TemplateError("「必要人数」シートの1行目にシフト名(早番・遅番など)がありません。")
    days, required = [], {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[0] in (None, ""):
            continue
        try:
            days.append(_to_date(row[0]))
        except ValueError as e:
            raise TemplateError(f"「必要人数」シートの日付「{row[0]}」を読み取れません。") from e
        d = len(days) - 1
        for t in range(len(shifts)):
            v = row[1 + t] if 1 + t < len(row) else 0
            required[(d, t)] = int(v) if v not in (None, "") else 0
    if not days:
        raise TemplateError("「必要人数」シートに日付の行がありません。")

    # --- 希望 ---
    ws = wb[SHEET_PREF]
    hdr = [c.value for c in ws[1]]
    col_of_day: dict[int, int] = {}
    day_index = {d: i for i, d in enumerate(days)}
    for ci, h in enumerate(hdr[1:], start=1):
        if h in (None, ""):
            continue
        try:
            dd = _to_date(h)
        except ValueError:
            continue
        if dd in day_index:
            col_of_day[ci] = day_index[dd]
    name_idx = {n: i for i, n in enumerate(names)}
    prefs = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[0] in (None, ""):
            continue
        nm = str(row[0]).strip()
        if nm not in name_idx:
            raise TemplateError(f"「希望」シートの「{nm}」が「スタッフ」シートにいません。")
        for ci, di in col_of_day.items():
            v = row[ci] if ci < len(row) else None
            v = "" if v is None else str(v).strip()
            if v in ("×", "x", "X", "✕"):
                prefs[(name_idx[nm], di)] = "×"
            elif v in ("△", "▲"):
                prefs[(name_idx[nm], di)] = "△"
    return staff, days, shifts, required, prefs


def build_shift_table(staff, days, shifts, assign):
    """表示用の2次元リスト(行=スタッフ, 列=日)を返す。空きは「」"""
    return [[shifts[assign[(s, d)]] if (s, d) in assign else "" for d in range(len(days))]
            for s in range(len(staff))]


def result_summary(staff, days, shifts, prefs, assign):
    """集計: スタッフ別勤務日数・希望反映率など。"""
    work = {st["name"]: 0 for st in staff}
    for (s, _d) in assign:
        work[staff[s]["name"]] += 1
    n_off = sum(1 for v in prefs.values() if v == "×")
    n_tri = sum(1 for v in prefs.values() if v == "△")
    tri_assigned = sum(1 for k, v in prefs.items() if v == "△" and k in assign)
    off_violate = sum(1 for k, v in prefs.items() if v == "×" and k in assign)
    total_req = n_off + n_tri
    rate = 1.0 if total_req == 0 else (total_req - tri_assigned - off_violate) / total_req
    return {
        "work_days": work,
        "pref_rate": rate,
        "n_off": n_off, "n_tri": n_tri,
        "tri_assigned": tri_assigned, "off_violate": off_violate,
    }


def write_result(staff, days, shifts, prefs, assign) -> bytes:
    """結果Excel(シフト表+集計)をバイト列で返す。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "シフト表"
    ws.cell(1, 1, "名前").fill = HEAD_FILL
    ws.cell(1, 1).font = HEAD_FONT
    for d, day in enumerate(days):
        c = ws.cell(1, 2 + d, f"{day.month}/{day.day}({WEEKDAY[day.weekday()]})")
        c.fill = HEAD_FILL
        c.font = HEAD_FONT
        c.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(2 + d)].width = 9
    ws.cell(1, 2 + len(days), "勤務日数").fill = HEAD_FILL
    ws.cell(1, 2 + len(days)).font = HEAD_FONT
    ws.column_dimensions["A"].width = 14

    table = build_shift_table(staff, days, shifts, assign)
    for s, st in enumerate(staff):
        ws.cell(2 + s, 1, st["name"] + ("★" if st["is_leader"] else ""))
        n = 0
        for d in range(len(days)):
            v = table[s][d]
            c = ws.cell(2 + s, 2 + d, v)
            c.alignment = Alignment(horizontal="center")
            if v:
                n += 1
                c.fill = PatternFill("solid", fgColor=SHIFT_COLORS[shifts.index(v) % len(SHIFT_COLORS)])
            elif prefs.get((s, d)) == "×":
                c.value = "休"
                c.font = Font(color="808080")
        ws.cell(2 + s, 2 + len(days), n)
    ws.freeze_panes = "B2"

    # 日別の人数行
    r0 = 3 + len(staff)
    for t, sh in enumerate(shifts):
        ws.cell(r0 + t, 1, f"{sh} 人数").font = Font(bold=True)
        for d in range(len(days)):
            ws.cell(r0 + t, 2 + d, sum(1 for s in range(len(staff)) if assign.get((s, d)) == t)).alignment = Alignment(horizontal="center")

    info = result_summary(staff, days, shifts, prefs, assign)
    ws2 = wb.create_sheet("集計")
    ws2.append(["項目", "値"])
    ws2.append(["希望反映率", f"{info['pref_rate']:.1%}"])
    ws2.append(["休み希望(×)の数", info["n_off"]])
    ws2.append(["△の数 / うち出勤させた数", f"{info['n_tri']} / {info['tri_assigned']}"])
    ws2.append([])
    ws2.append(["スタッフ", "勤務日数"])
    for k, v in info["work_days"].items():
        ws2.append([k, v])
    ws2.column_dimensions["A"].width = 26

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def make_template(path_or_buf, staff=None, days=None, shifts=None, required=None, prefs=None):
    """テンプレExcelを書き出す(引数省略時は空のひな形)。"""
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_STAFF
    ws.append(["名前", "週の上限日数", "リーダー可(◯)"])
    if staff is None:
        ws.append(["例)山田", 5, "◯"])
    else:
        for st in staff:
            ws.append([st["name"], st["max_days"], "◯" if st["is_leader"] else ""])

    if days is None:
        start = date.today().replace(day=1)
        days = [start + timedelta(days=i) for i in range(30)]
    shifts = shifts or ["早番", "遅番"]
    if required is None:
        required = {(d, t): 2 for d in range(len(days)) for t in range(len(shifts))}

    wp = wb.create_sheet(SHEET_PREF)
    wp.append(["名前"] + days)
    names = [s["name"] for s in staff] if staff else ["例)山田"]
    for si, nm in enumerate(names):
        wp.append([nm] + [(prefs or {}).get((si, d), "") for d in range(len(days))])

    wr = wb.create_sheet(SHEET_REQ)
    wr.append(["日付"] + shifts)
    for d, day in enumerate(days):
        wr.append([day] + [required.get((d, t), 0) for t in range(len(shifts))])

    for w in (ws, wp, wr):
        for c in w[1]:
            c.fill = HEAD_FILL
            c.font = HEAD_FONT
            if isinstance(c.value, (date, datetime)):
                c.number_format = "m/d"
        w.column_dimensions["A"].width = 14
    for row in wr.iter_rows(min_row=2, max_col=1):
        row[0].number_format = "yyyy/mm/dd"
    wb.save(path_or_buf)


def sample_data(seed: int = 7, n_staff: int = 10, n_days: int = 30, start: date | None = None):
    """飲食店想定のサンプル: 10人×30日、早番・遅番。"""
    rnd = random.Random(seed)
    start = start or date(2026, 11, 1)
    days = [start + timedelta(days=i) for i in range(n_days)]
    names = ["佐藤", "鈴木", "高橋", "田中", "伊藤", "渡辺", "山本", "中村", "小林", "加藤"][:n_staff]
    staff = []
    for i, nm in enumerate(names):
        staff.append({"name": nm, "max_days": 5 if i < 7 else 4, "is_leader": i in (0, 1, 2, 3)})
    shifts = ["早番", "遅番"]
    required = {}
    for d, day in enumerate(days):
        weekend = day.weekday() >= 5
        required[(d, 0)] = 3 if weekend else 2
        required[(d, 1)] = 3 if weekend else 2
    prefs = {}
    for s in range(n_staff):
        for d in range(n_days):
            r = rnd.random()
            if r < 0.15:
                prefs[(s, d)] = "×"
            elif r < 0.25:
                prefs[(s, d)] = "△"
    # デモが必ず解ける(現実的な)データになるよう、各日に「必要人数+2人」以上は
    # 出勤可能(×でない)になるまで、×を△に置き換える
    for d in range(n_days):
        need = required[(d, 0)] + required[(d, 1)]
        offs = [s for s in range(n_staff) if prefs.get((s, d)) == "×"]
        rnd.shuffle(offs)
        while n_staff - len(offs) < need + 2 and offs:
            prefs[(offs.pop(), d)] = "△"
    return staff, days, shifts, required, prefs
